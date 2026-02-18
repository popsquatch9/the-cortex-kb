from __future__ import annotations

import hashlib
import json
import os
import re
import sqlite3
import textwrap
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Optional
from urllib.parse import parse_qs, unquote, urlparse

import typer
from rich.console import Console
from rich.markdown import Markdown as RichMarkdown
from rich.panel import Panel
from rich.prompt import Prompt
from rich.table import Table
from rich import box

try:
    import chromadb

    HAS_CHROMA = True
except ImportError:
    HAS_CHROMA = False

try:
    from sentence_transformers import SentenceTransformer

    HAS_SBERT = True
except ImportError:
    HAS_SBERT = False

try:
    import requests
    from bs4 import BeautifulSoup

    HAS_WEB = True
except ImportError:
    HAS_WEB = False

try:
    import PyPDF2

    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    import docx

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import feedparser

    HAS_FEED = True
except ImportError:
    HAS_FEED = False

try:
    import openai

    HAS_OPENAI = True
except ImportError:
    HAS_OPENAI = False

try:
    import schedule

    HAS_SCHEDULE = True
except ImportError:
    HAS_SCHEDULE = False


CORTEX_DIR = Path.home() / ".cortex"
DB_PATH = CORTEX_DIR / "cortex.db"
CHROMA_DIR = CORTEX_DIR / "vectors"
SOURCES_DIR = CORTEX_DIR / "sources"
EXPORT_DIR = CORTEX_DIR / "exports"
OBSIDIAN_DIR = CORTEX_DIR / "obsidian_vault"
EMBED_MODEL = "all-MiniLM-L6-v2"
CHUNK_SIZE = 512
CHUNK_OVERLAP = 64
QUALITY_THRESHOLD = 0.55
MAX_EXPAND_RESULTS = 8

FEED_PULL_INTERVAL = 2
EXPAND_INTERVAL = 12

console = Console()
app = typer.Typer(help="The Cortex — Your Personal Knowledge Base")
feeds_app = typer.Typer(help="Manage RSS/Atom feed subscriptions")
export_app = typer.Typer(help="Export knowledge to various formats")
app.add_typer(feeds_app, name="feeds")
app.add_typer(export_app, name="export")


class SourceType(str, Enum):
    FILE = "file"
    URL = "url"
    NOTE = "note"
    AUTO = "auto-discovered"
    RSS = "rss-feed"


class QualityTier(str, Enum):
    GOLD = "gold"
    SILVER = "silver"
    BRONZE = "bronze"
    REJECTED = "rejected"


@dataclass
class KnowledgeEntry:
    id: str
    title: str
    source: str
    source_type: SourceType
    content: str
    summary: str = ""
    tags: list[str] = field(default_factory=list)
    quality: QualityTier = QualityTier.GOLD
    created_at: str = field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    parent_id: Optional[str] = None
    metadata: dict = field(default_factory=dict)

    def fingerprint(self) -> str:
        return hashlib.sha256(self.content[:2000].encode()).hexdigest()[:16]


class Cortex:
    """The brain. Reads, reasons, organizes, expands, converses."""

    def __init__(self):
        for d in [CORTEX_DIR, SOURCES_DIR, EXPORT_DIR, OBSIDIAN_DIR]:
            d.mkdir(exist_ok=True)
        self._init_sqlite()
        self._init_vectors()
        self._init_embedder()

    def _init_sqlite(self):
        self.db = sqlite3.connect(str(DB_PATH), check_same_thread=False)
        self.db.row_factory = sqlite3.Row
        self.db.executescript(
            """
            CREATE TABLE IF NOT EXISTS entries (
                id TEXT PRIMARY KEY,
                title TEXT,
                source TEXT,
                source_type TEXT,
                content TEXT,
                summary TEXT,
                quality TEXT DEFAULT 'gold',
                created_at TEXT,
                parent_id TEXT,
                metadata TEXT DEFAULT '{}'
            );
            CREATE TABLE IF NOT EXISTS tags (
                entry_id TEXT,
                tag TEXT,
                FOREIGN KEY (entry_id) REFERENCES entries(id),
                UNIQUE(entry_id, tag)
            );
            CREATE TABLE IF NOT EXISTS relationships (
                source_id TEXT,
                target_id TEXT,
                relation TEXT,
                score REAL,
                FOREIGN KEY (source_id) REFERENCES entries(id),
                FOREIGN KEY (target_id) REFERENCES entries(id)
            );
            CREATE TABLE IF NOT EXISTS feeds (
                id TEXT PRIMARY KEY,
                url TEXT UNIQUE,
                name TEXT,
                tag TEXT DEFAULT '',
                last_pulled TEXT,
                added_at TEXT,
                active INTEGER DEFAULT 1
            );
            CREATE TABLE IF NOT EXISTS feed_seen (
                feed_id TEXT,
                article_hash TEXT,
                UNIQUE(feed_id, article_hash)
            );
            CREATE TABLE IF NOT EXISTS chat_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                role TEXT,
                content TEXT,
                timestamp TEXT
            );
            CREATE INDEX IF NOT EXISTS idx_tags ON tags(tag);
            CREATE INDEX IF NOT EXISTS idx_quality ON entries(quality);
            CREATE INDEX IF NOT EXISTS idx_feeds_active ON feeds(active);
        """
        )
        self.db.commit()

    def _init_vectors(self):
        if HAS_CHROMA:
            self.vdb = chromadb.PersistentClient(path=str(CHROMA_DIR))
            self.collection = self.vdb.get_or_create_collection(
                name="cortex",
                metadata={"hnsw:space": "cosine"},
            )
        else:
            self.collection = None
            console.print("[yellow]⚠ ChromaDB not installed — semantic search disabled[/]")

    def _init_embedder(self):
        if HAS_SBERT:
            self.embedder = SentenceTransformer(EMBED_MODEL)
        else:
            self.embedder = None
            console.print("[yellow]⚠ sentence-transformers not installed — keyword search only[/]")

    @staticmethod
    def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
        words = text.split()
        chunks, i = [], 0
        while i < len(words):
            chunk = " ".join(words[i : i + size])
            if chunk.strip():
                chunks.append(chunk)
            i += size - overlap
        return chunks or [text]

    @staticmethod
    def extract_keywords(text: str, top_n: int = 12) -> list[str]:
        stop = {
            "the",
            "a",
            "an",
            "is",
            "are",
            "was",
            "were",
            "be",
            "been",
            "being",
            "have",
            "has",
            "had",
            "do",
            "does",
            "did",
            "will",
            "would",
            "could",
            "should",
            "may",
            "might",
            "shall",
            "can",
            "need",
            "dare",
            "ought",
            "used",
            "to",
            "of",
            "in",
            "for",
            "on",
            "with",
            "at",
            "by",
            "from",
            "as",
            "into",
            "through",
            "during",
            "before",
            "after",
            "above",
            "below",
            "between",
            "out",
            "off",
            "over",
            "under",
            "again",
            "further",
            "then",
            "once",
            "here",
            "there",
            "when",
            "where",
            "why",
            "how",
            "all",
            "each",
            "every",
            "both",
            "few",
            "more",
            "most",
            "other",
            "some",
            "such",
            "no",
            "nor",
            "not",
            "only",
            "own",
            "same",
            "so",
            "than",
            "too",
            "very",
            "just",
            "because",
            "but",
            "and",
            "or",
            "if",
            "while",
            "this",
            "that",
            "these",
            "those",
            "it",
            "its",
            "i",
            "me",
            "my",
            "we",
            "our",
            "you",
            "your",
            "he",
            "him",
            "his",
            "she",
            "her",
            "they",
            "them",
            "their",
            "what",
            "which",
            "who",
            "whom",
            "about",
            "also",
            "new",
            "one",
            "two",
        }
        words = re.findall(r"\b[a-z]{3,}\b", text.lower())
        freq = {}
        for w in words:
            if w not in stop:
                freq[w] = freq.get(w, 0) + 1
        return sorted(freq, key=freq.get, reverse=True)[:top_n]

    def auto_tag(self, text: str) -> list[str]:
        keywords = self.extract_keywords(text, top_n=8)
        domain_signals = {
            "health": [
                "herb",
                "treatment",
                "symptom",
                "diagnosis",
                "medicine",
                "ayurveda",
                "tcm",
                "homeopathy",
                "autoimmune",
                "inflammation",
                "lupus",
                "endometriosis",
                "pcos",
                "perimenopause",
            ],
            "cybersecurity": [
                "vulnerability",
                "encryption",
                "firewall",
                "malware",
                "network",
                "security",
                "threat",
                "attack",
                "protocol",
            ],
            "data-science": [
                "dataset",
                "analysis",
                "regression",
                "model",
                "statistics",
                "visualization",
                "pandas",
                "python",
                "correlation",
            ],
            "web3": [
                "blockchain",
                "solidity",
                "ethereum",
                "smart-contract",
                "defi",
                "token",
                "wallet",
                "crypto",
                "decentralized",
            ],
            "education": [
                "student",
                "learning",
                "curriculum",
                "iep",
                "special-education",
                "teaching",
                "assessment",
                "accommodation",
                "paraprofessional",
            ],
            "psychedelics": [
                "psilocybin",
                "mdma",
                "ketamine",
                "maps",
                "therapy",
                "psychedelic",
                "integration",
                "ceremony",
            ],
            "philosophy": [
                "consciousness",
                "existence",
                "meaning",
                "ontology",
                "epistemology",
                "metaphysics",
                "phenomenology",
            ],
            "herbalism": [
                "tincture",
                "decoction",
                "adaptogen",
                "nervine",
                "herb",
                "botanical",
                "extract",
                "infusion",
                "ashwagandha",
            ],
            "death-work": [
                "death",
                "doula",
                "dying",
                "grief",
                "end-of-life",
                "palliative",
                "hospice",
                "bereavement",
                "mortality",
            ],
            "yoga-meditation": [
                "yoga",
                "meditation",
                "pranayama",
                "asana",
                "mindfulness",
                "chakra",
                "mantra",
                "breathwork",
                "samadhi",
            ],
        }
        tags = set()
        text_lower = text.lower()
        for domain, signals in domain_signals.items():
            if any(s in text_lower for s in signals):
                tags.add(domain)
        for kw in keywords[:5]:
            tags.add(kw)
        return sorted(tags)

    def summarize(self, text: str, max_len: int = 300) -> str:
        if HAS_OPENAI and os.getenv("OPENAI_API_KEY"):
            try:
                client = openai.OpenAI()
                resp = client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[
                        {
                            "role": "system",
                            "content": "Summarize concisely. Preserve key facts, data, and nuance.",
                        },
                        {"role": "user", "content": text[:4000]},
                    ],
                    max_tokens=200,
                )
                return resp.choices[0].message.content.strip()
            except Exception:
                pass

        sentences = re.split(r"(?<=[.!?])\s+", text.strip())
        summary = ""
        for s in sentences:
            if len(summary) + len(s) < max_len:
                summary += s + " "
            else:
                break
        return summary.strip() or text[:max_len]

    def _read_file(self, path: Path) -> tuple[str, str]:
        suffix = path.suffix.lower()
        if suffix == ".pdf" and HAS_PDF:
            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                text = "\n".join(p.extract_text() or "" for p in reader.pages)
            return path.stem, text
        if suffix == ".docx" and HAS_DOCX:
            doc = docx.Document(str(path))
            return path.stem, "\n".join(p.text for p in doc.paragraphs)

        try:
            return path.stem, path.read_text(encoding="utf-8", errors="ignore")
        except Exception as e:
            raise ValueError(f"Can't read {path}: {e}")

    def _read_url(self, url: str) -> tuple[str, str]:
        if not HAS_WEB:
            raise ImportError("Install requests + beautifulsoup4")
        headers = {"User-Agent": "Mozilla/5.0 (Cortex Knowledge Bot) AppleWebKit/537.36"}
        resp = requests.get(url, headers=headers, timeout=15)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header", "aside", "iframe"]):
            tag.decompose()
        title = soup.title.string.strip() if soup.title and soup.title.string else urlparse(url).netloc
        main = soup.find("article") or soup.find("main") or soup.find("body")
        text = main.get_text(separator="\n", strip=True) if main else soup.get_text(separator="\n", strip=True)
        return title, re.sub(r"\n{3,}", "\n\n", text)

    def ingest(
        self,
        source: str,
        source_type: SourceType = None,
        parent_id: str = None,
        extra_tags: list[str] = None,
        quality: QualityTier = None,
        quiet: bool = False,
    ) -> Optional[KnowledgeEntry]:
        if source_type is None:
            if source.startswith(("http://", "https://")):
                source_type = SourceType.URL
            elif Path(source).exists():
                source_type = SourceType.FILE
            else:
                source_type = SourceType.NOTE

        try:
            if source_type in (SourceType.URL, SourceType.AUTO, SourceType.RSS) and source.startswith(
                ("http://", "https://")
            ):
                title, content = self._read_url(source)
            elif source_type == SourceType.FILE:
                title, content = self._read_file(Path(source))
            else:
                title, content = source[:60], source
        except Exception as e:
            if not quiet:
                console.print(f"[red]✗ Failed to read source: {e}[/]")
            return None

        if not content or len(content.strip()) < 20:
            if not quiet:
                console.print("[red]✗ Content too short or empty — skipping.[/]")
            return None

        dup = self.db.execute(
            "SELECT id FROM entries WHERE substr(content, 1, 2000) = ?",
            (content[:2000],),
        ).fetchone()
        if dup:
            if not quiet:
                console.print(f"[yellow]⚠ Duplicate detected — skipping '{title[:50]}'[/]")
            return None

        entry_id = hashlib.sha256(f"{source}{time.time()}".encode()).hexdigest()[:12]
        all_tags = self.auto_tag(content)
        if extra_tags:
            all_tags = sorted(set(all_tags + extra_tags))

        if quality is None:
            quality = (
                QualityTier.GOLD
                if parent_id is None and source_type not in (SourceType.AUTO, SourceType.RSS)
                else QualityTier.SILVER
            )

        entry = KnowledgeEntry(
            id=entry_id,
            title=title,
            source=source,
            source_type=source_type,
            content=content,
            summary=self.summarize(content),
            tags=all_tags,
            quality=quality,
            parent_id=parent_id,
        )

        self.db.execute(
            "INSERT INTO entries (id,title,source,source_type,content,summary,quality,created_at,parent_id,metadata) VALUES (?,?,?,?,?,?,?,?,?,?)",
            (
                entry.id,
                entry.title,
                entry.source,
                entry.source_type.value,
                entry.content,
                entry.summary,
                entry.quality.value,
                entry.created_at,
                entry.parent_id,
                json.dumps(entry.metadata),
            ),
        )
        for tag in entry.tags:
            self.db.execute("INSERT OR IGNORE INTO tags (entry_id, tag) VALUES (?,?)", (entry.id, tag))
        self.db.commit()

        if self.collection and self.embedder:
            chunks = self.chunk_text(content)
            embeddings = self.embedder.encode(chunks).tolist()
            ids = [f"{entry.id}_chunk_{i}" for i in range(len(chunks))]
            metadatas = [{"entry_id": entry.id, "title": title, "chunk_idx": i} for i in range(len(chunks))]
            self.collection.add(ids=ids, embeddings=embeddings, documents=chunks, metadatas=metadatas)

        self._link_related(entry)

        if not quiet:
            console.print(
                Panel(
                    f"[bold green]✓ Ingested:[/] {title}\n"
                    f"[dim]Tags:[/] {', '.join(entry.tags)}\n"
                    f"[dim]Quality:[/] {entry.quality.value}\n"
                    f"[dim]Chunks:[/] {len(self.chunk_text(content))}",
                    title="📥 Knowledge Added",
                    border_style="green",
                )
            )
        return entry

    def _link_related(self, entry: KnowledgeEntry):
        if not self.collection or not self.embedder:
            return
        query_emb = self.embedder.encode([entry.summary]).tolist()
        try:
            results = self.collection.query(query_embeddings=query_emb, n_results=5)
        except Exception:
            return
        seen = set()
        for meta, dist in zip(results["metadatas"][0], results["distances"][0]):
            target_id = meta["entry_id"]
            if target_id == entry.id or target_id in seen:
                continue
            seen.add(target_id)
            score = 1 - dist
            if score > 0.3:
                self.db.execute(
                    "INSERT INTO relationships (source_id,target_id,relation,score) VALUES (?,?,?,?)",
                    (entry.id, target_id, "related", round(score, 4)),
                )
        self.db.commit()

    def search(self, query: str, top_k: int = 8) -> list[dict]:
        results = []
        if self.collection and self.embedder:
            query_emb = self.embedder.encode([query]).tolist()
            vresults = self.collection.query(query_embeddings=query_emb, n_results=top_k * 2)
            seen_entries = set()
            for doc, meta, dist in zip(
                vresults["documents"][0],
                vresults["metadatas"][0],
                vresults["distances"][0],
            ):
                eid = meta["entry_id"]
                if eid in seen_entries:
                    continue
                seen_entries.add(eid)
                row = self.db.execute("SELECT * FROM entries WHERE id = ?", (eid,)).fetchone()
                if row:
                    tags = [
                        r["tag"]
                        for r in self.db.execute("SELECT tag FROM tags WHERE entry_id = ?", (eid,)).fetchall()
                    ]
                    results.append(
                        {
                            "id": row["id"],
                            "title": row["title"],
                            "summary": row["summary"],
                            "source": row["source"],
                            "quality": row["quality"],
                            "tags": tags,
                            "relevance": round(1 - dist, 4),
                            "snippet": doc[:300],
                        }
                    )

        keywords = self.extract_keywords(query, top_n=5)
        if keywords:
            placeholders = " OR ".join(["content LIKE ?" for _ in keywords])
            params = [f"%{kw}%" for kw in keywords]
            rows = self.db.execute(f"SELECT * FROM entries WHERE {placeholders} LIMIT ?", params + [top_k]).fetchall()
            existing_ids = {r["id"] for r in results}
            for row in rows:
                if row["id"] not in existing_ids:
                    tags = [
                        r["tag"]
                        for r in self.db.execute(
                            "SELECT tag FROM tags WHERE entry_id = ?", (row["id"],)
                        ).fetchall()
                    ]
                    results.append(
                        {
                            "id": row["id"],
                            "title": row["title"],
                            "summary": row["summary"],
                            "source": row["source"],
                            "quality": row["quality"],
                            "tags": tags,
                            "relevance": 0.5,
                            "snippet": row["content"][:300],
                        }
                    )

        results.sort(key=lambda x: x["relevance"], reverse=True)
        return results[:top_k]

    def _search_web(self, query: str, num: int = 5) -> list[str]:
        if not HAS_WEB:
            return []
        urls = []

        def normalize_result_url(href: str) -> str:
            if href.startswith("http"):
                return href
            if href.startswith("/"):
                parsed = urlparse(href)
                params = parse_qs(parsed.query)
                uddg = params.get("uddg", [""])[0]
                if uddg:
                    return unquote(uddg)
            return ""

        try:
            resp = requests.get(
                "https://html.duckduckgo.com/html/",
                params={"q": query},
                headers={"User-Agent": "Mozilla/5.0"},
                timeout=10,
            )
            soup = BeautifulSoup(resp.text, "html.parser")
            for a in soup.select(".result__a"):
                href = normalize_result_url(a.get("href", ""))
                if href.startswith("http") and href not in urls:
                    urls.append(href)
                if len(urls) >= num:
                    break
        except Exception:
            pass
        return urls

    def _search_scholar(self, query: str, num: int = 3) -> list[str]:
        if not HAS_WEB:
            return []
        urls = []
        try:
            resp = requests.get(
                "https://scholar.google.com/scholar",
                params={"q": query, "num": num},
                headers={"User-Agent": "Mozilla/5.0"},
                timeout=10,
            )
            soup = BeautifulSoup(resp.text, "html.parser")
            for h3 in soup.select(".gs_rt a"):
                href = h3.get("href", "")
                if href.startswith("http"):
                    urls.append(href)
        except Exception:
            pass
        return urls[:num]

    def _search_wikipedia(self, query: str, num: int = 4) -> list[str]:
        if not HAS_WEB:
            return []
        try:
            resp = requests.get(
                "https://en.wikipedia.org/w/api.php",
                params={
                    "action": "query",
                    "format": "json",
                    "list": "search",
                    "srsearch": query,
                    "srlimit": num,
                },
                headers={"User-Agent": "Mozilla/5.0"},
                timeout=10,
            )
            resp.raise_for_status()
            data = resp.json()
            urls = []
            for item in data.get("query", {}).get("search", []):
                title = item.get("title", "").replace(" ", "_")
                if title:
                    urls.append(f"https://en.wikipedia.org/wiki/{title}")
            return urls[:num]
        except Exception:
            return []

    def _vet_source(self, url: str, reference_text: str) -> tuple[bool, float]:
        try:
            _, content = self._read_url(url)
        except Exception:
            return False, 0.0
        if len(content.strip()) < 100:
            return False, 0.0

        if self.embedder:
            emb_ref = self.embedder.encode([reference_text[:1000]])
            emb_new = self.embedder.encode([content[:1000]])
            from numpy import dot
            from numpy.linalg import norm

            sim = float(dot(emb_ref[0], emb_new[0]) / (norm(emb_ref[0]) * norm(emb_new[0])))
        else:
            ref_kw = set(self.extract_keywords(reference_text))
            new_kw = set(self.extract_keywords(content))
            sim = len(ref_kw & new_kw) / max(len(ref_kw | new_kw), 1)

        trusted = [
            "nih.gov",
            "pubmed",
            "scholar.google",
            "nature.com",
            "sciencedirect",
            "arxiv.org",
            "wikipedia.org",
            "edu",
            "ncbi.nlm.nih.gov",
            "who.int",
            "mayoclinic.org",
        ]
        domain = urlparse(url).netloc.lower()
        trust_bonus = 0.1 if any(td in domain for td in trusted) else 0.0
        final_score = sim + trust_bonus
        return final_score >= QUALITY_THRESHOLD, round(final_score, 4)

    def expand(self, topic: str = None, quiet: bool = False):
        if not quiet:
            console.print(Panel("🔍 [bold]Expanding knowledge base...[/]", border_style="cyan"))

        if topic:
            queries = [topic]
        else:
            top_tags = self.db.execute(
                "SELECT tag, COUNT(*) as cnt FROM tags GROUP BY tag ORDER BY cnt DESC LIMIT 5"
            ).fetchall()
            recent = self.db.execute("SELECT title, summary FROM entries ORDER BY created_at DESC LIMIT 3").fetchall()
            queries = [t["tag"] for t in top_tags]
            for r in recent:
                queries.append(f"{r['title']} research")

        if not queries:
            if not quiet:
                console.print("[yellow]Nothing to expand on yet — ingest some knowledge first![/]")
            return

        existing_sources = {r["source"] for r in self.db.execute("SELECT source FROM entries").fetchall()}
        added = 0

        for query in queries[:5]:
            if not quiet:
                console.print(f"  🔎 Searching: [cyan]{query}[/]")
            urls = self._search_web(query + " research evidence", num=4)
            urls += self._search_scholar(query, num=2)
            if len(urls) < 2:
                urls += self._search_wikipedia(query, num=4)

            deduped = []
            seen = set()
            for u in urls:
                if u and u not in seen:
                    seen.add(u)
                    deduped.append(u)
            urls = deduped

            for url in urls:
                if url in existing_sources or added >= MAX_EXPAND_RESULTS:
                    continue
                ref_rows = self.db.execute(
                    "SELECT content FROM entries WHERE id IN (SELECT entry_id FROM tags WHERE tag LIKE ?) LIMIT 1",
                    (f"%{query.split()[0]}%",),
                ).fetchall()
                ref_text = ref_rows[0]["content"] if ref_rows else query
                passed, score = self._vet_source(url, ref_text)

                if passed:
                    if not quiet:
                        console.print(f"    [green]✓[/] Vetted ({score:.2f}): {url[:80]}")
                    entry = self.ingest(url, SourceType.AUTO, quiet=quiet)
                    if entry:
                        added += 1
                elif not quiet:
                    console.print(f"    [dim]✗ Rejected ({score:.2f}): {url[:80]}[/]")

        if not quiet:
            console.print(f"\n[bold green]Expansion complete: {added} new sources added.[/]")

    def feed_add(self, url: str, name: str = None, tag: str = ""):
        if not HAS_FEED:
            console.print("[red]Install feedparser: pip install feedparser[/]")
            return
        feed = feedparser.parse(url)
        if not feed.entries:
            console.print(f"[red]✗ No entries found at {url} — is this a valid RSS/Atom feed?[/]")
            return
        feed_name = name or feed.feed.get("title", urlparse(url).netloc)
        feed_id = hashlib.sha256(url.encode()).hexdigest()[:10]
        try:
            self.db.execute(
                "INSERT INTO feeds (id,url,name,tag,added_at,active) VALUES (?,?,?,?,?,1)",
                (feed_id, url, feed_name, tag, datetime.now(timezone.utc).isoformat()),
            )
            self.db.commit()
            console.print(
                Panel(
                    f"[bold green]✓ Feed added:[/] {feed_name}\n"
                    f"[dim]URL:[/] {url}\n"
                    f"[dim]Tag:[/] {tag or '(auto)'}\n"
                    f"[dim]Articles available:[/] {len(feed.entries)}",
                    title="📡 Feed Subscribed",
                    border_style="green",
                )
            )
        except sqlite3.IntegrityError:
            console.print(f"[yellow]⚠ Feed already exists: {url}[/]")

    def feed_list(self):
        feeds = self.db.execute("SELECT * FROM feeds WHERE active = 1 ORDER BY name").fetchall()
        if not feeds:
            console.print("[yellow]No feeds subscribed yet. Use: cortex.py feeds add <url>[/]")
            return
        table = Table(title="📡 Subscribed Feeds", box=box.DOUBLE_EDGE, border_style="cyan")
        table.add_column("ID", style="dim", width=10)
        table.add_column("Name", style="bold cyan")
        table.add_column("Tag", style="green")
        table.add_column("Last Pulled", style="dim")
        table.add_column("URL", style="dim", max_width=40)
        for f in feeds:
            last = f["last_pulled"][:16] if f["last_pulled"] else "never"
            table.add_row(f["id"], f["name"], f["tag"] or "(auto)", last, f["url"][:40])
        console.print(table)

    def feed_remove(self, feed_id: str):
        self.db.execute("UPDATE feeds SET active = 0 WHERE id = ?", (feed_id,))
        self.db.commit()
        console.print(f"[green]✓ Feed {feed_id} deactivated.[/]")

    def feed_pull(self, quiet: bool = False):
        if not HAS_FEED:
            console.print("[red]Install feedparser: pip install feedparser[/]")
            return
        feeds = self.db.execute("SELECT * FROM feeds WHERE active = 1").fetchall()
        if not feeds:
            if not quiet:
                console.print("[yellow]No active feeds.[/]")
            return

        total_new = 0
        for f in feeds:
            if not quiet:
                console.print(f"  📡 Pulling: [cyan]{f['name']}[/]")
            feed = feedparser.parse(f["url"])
            new_count = 0

            for entry in feed.entries[:10]:
                article_hash = hashlib.sha256((entry.get("link", "") + entry.get("title", "")).encode()).hexdigest()[:16]
                seen = self.db.execute(
                    "SELECT 1 FROM feed_seen WHERE feed_id = ? AND article_hash = ?",
                    (f["id"], article_hash),
                ).fetchone()
                if seen:
                    continue
                self.db.execute(
                    "INSERT OR IGNORE INTO feed_seen (feed_id, article_hash) VALUES (?,?)",
                    (f["id"], article_hash),
                )
                url = entry.get("link", "")
                if url:
                    extra_tags = [f["tag"]] if f["tag"] else []
                    result = self.ingest(
                        url,
                        SourceType.RSS,
                        extra_tags=extra_tags,
                        quality=QualityTier.SILVER,
                        quiet=quiet,
                    )
                    if result:
                        new_count += 1

            self.db.execute(
                "UPDATE feeds SET last_pulled = ? WHERE id = ?",
                (datetime.now(timezone.utc).isoformat(), f["id"]),
            )
            self.db.commit()
            total_new += new_count
            if not quiet:
                console.print(f"    → {new_count} new articles ingested")

        if not quiet:
            console.print(f"\n[bold green]Feed pull complete: {total_new} new articles total.[/]")

    def export_obsidian(self, vault_dir: Path = None):
        vault = vault_dir or OBSIDIAN_DIR
        vault.mkdir(parents=True, exist_ok=True)

        entries = self.db.execute("SELECT * FROM entries ORDER BY created_at DESC").fetchall()
        if not entries:
            console.print("[yellow]Nothing to export yet.[/]")
            return

        all_tags = self.get_tags()
        tags_dir = vault / "Tags"
        tags_dir.mkdir(exist_ok=True)
        for q in ["Gold", "Silver", "Bronze"]:
            (vault / q).mkdir(exist_ok=True)

        exported = 0
        for entry in entries:
            tags = [
                r["tag"]
                for r in self.db.execute("SELECT tag FROM tags WHERE entry_id = ?", (entry["id"],)).fetchall()
            ]
            relations = self.db.execute(
                """
                SELECT e.title, e.id, r.relation, r.score
                FROM relationships r JOIN entries e ON e.id = r.target_id
                WHERE r.source_id = ? ORDER BY r.score DESC LIMIT 10
                """,
                (entry["id"],),
            ).fetchall()

            safe_title = re.sub(r'[<>:"/\\|?*]', "_", entry["title"][:80])
            quality_folder = entry["quality"].capitalize()

            md = "---\n"
            md += f"id: {entry['id']}\n"
            md += f"source: \"{entry['source']}\"\n"
            md += f"type: {entry['source_type']}\n"
            md += f"quality: {entry['quality']}\n"
            md += f"created: {entry['created_at'][:10]}\n"
            md += "tags:\n"
            for t in tags:
                md += f"  - {t}\n"
            md += "---\n\n"
            md += f"# {entry['title']}\n\n"
            md += f"## Summary\n{entry['summary']}\n\n"

            if relations:
                md += "## Related Knowledge\n"
                for rel in relations:
                    rel_title = re.sub(r'[<>:"/\\|?*]', "_", rel["title"][:80])
                    md += f"- [[{rel_title}]] ({rel['relation']}, {rel['score']:.0%})\n"
                md += "\n"

            if tags:
                md += "## Tags\n"
                md += " ".join(f"#{t}" for t in tags) + "\n\n"

            md += f"## Source\n{entry['source']}\n\n"
            md += f"## Full Content\n{entry['content'][:5000]}\n"

            filepath = vault / quality_folder / f"{safe_title}.md"
            filepath.write_text(md, encoding="utf-8")
            exported += 1

        for tag, count in all_tags.items():
            tag_entries = self.db.execute(
                """
                SELECT e.title, e.id, e.quality, e.summary
                FROM entries e JOIN tags t ON e.id = t.entry_id
                WHERE t.tag = ? ORDER BY e.created_at DESC
                """,
                (tag,),
            ).fetchall()

            moc = f"# {tag.title()} — Map of Content\n\n"
            moc += f"*{count} entries*\n\n"
            for te in tag_entries:
                safe = re.sub(r'[<>:"/\\|?*]', "_", te["title"][:80])
                moc += f"- [[{safe}]] ({te['quality']}) — {te['summary'][:100]}\n"

            safe_tag = re.sub(r'[<>:"/\\|?*]', "_", tag)
            (tags_dir / f"MOC — {safe_tag}.md").write_text(moc, encoding="utf-8")

        index = "# 🧠 The Cortex — Knowledge Index\n\n"
        index += (
            f"*{len(entries)} entries · {len(all_tags)} tags · "
            f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}*\n\n"
        )
        index += "## Tag Maps of Content\n"
        for tag in sorted(all_tags):
            safe_tag = re.sub(r'[<>:"/\\|?*]', "_", tag)
            index += f"- [[MOC — {safe_tag}]] ({all_tags[tag]})\n"
        index += "\n## Recent Entries\n"
        for entry in entries[:20]:
            safe = re.sub(r'[<>:"/\\|?*]', "_", entry["title"][:80])
            index += f"- [[{safe}]] — {entry['created_at'][:10]}\n"

        (vault / "INDEX.md").write_text(index, encoding="utf-8")

        console.print(
            Panel(
                f"[bold green]✓ Exported {exported} entries to Obsidian vault[/]\n"
                f"[dim]Location:[/] {vault}\n"
                f"[dim]Tag MOCs:[/] {len(all_tags)}\n"
                f"[dim]Master index:[/] INDEX.md",
                title="📓 Obsidian Export",
                border_style="green",
            )
        )

    def export_json(self, filepath: Path = None):
        filepath = filepath or (EXPORT_DIR / f"cortex_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json")
        entries = self.db.execute("SELECT * FROM entries ORDER BY created_at DESC").fetchall()
        data = {
            "exported_at": datetime.now(timezone.utc).isoformat(),
            "stats": self.get_stats(),
            "entries": [],
        }
        for entry in entries:
            tags = [
                r["tag"]
                for r in self.db.execute("SELECT tag FROM tags WHERE entry_id = ?", (entry["id"],)).fetchall()
            ]
            relations = [
                dict(r)
                for r in self.db.execute(
                    "SELECT target_id, relation, score FROM relationships WHERE source_id = ?",
                    (entry["id"],),
                ).fetchall()
            ]
            data["entries"].append(
                {
                    "id": entry["id"],
                    "title": entry["title"],
                    "source": entry["source"],
                    "source_type": entry["source_type"],
                    "summary": entry["summary"],
                    "quality": entry["quality"],
                    "tags": tags,
                    "relationships": relations,
                    "created_at": entry["created_at"],
                    "content_preview": entry["content"][:500],
                }
            )
        filepath.parent.mkdir(parents=True, exist_ok=True)
        filepath.write_text(json.dumps(data, indent=2), encoding="utf-8")
        console.print(f"[green]✓ Exported to {filepath} ({len(entries)} entries)[/]")

    def chat(self):
        console.print(
            Panel(
                "[bold magenta]🧠 Cortex Chat[/]\n"
                "[dim]Ask me anything about your knowledge base.\n"
                "I'll search, synthesize, and reason across everything you've stored.\n\n"
                "Commands:  /search <query>  /expand <topic>  /ingest <source>\n"
                "           /tags  /stats  /help  /quit[/]",
                border_style="magenta",
            )
        )

        history = []

        while True:
            try:
                user_input = Prompt.ask("\n[bold cyan]You[/]")
            except (KeyboardInterrupt, EOFError):
                console.print("\n[dim]Later. Knowledge persists. 🧠[/]")
                break

            if not user_input.strip():
                continue

            if user_input.strip().startswith("/"):
                self._handle_slash_command(user_input.strip())
                continue

            if user_input.strip().lower() in ("quit", "exit", "/quit", "/exit"):
                console.print("[dim]Later. Knowledge persists. 🧠[/]")
                break

            results = self.search(user_input, top_k=5)
            context_parts, sources_used = [], []

            for r in results:
                if r["relevance"] > 0.35:
                    context_parts.append(
                        f"[{r['quality'].upper()}] {r['title']}\n"
                        f"Tags: {', '.join(r['tags'][:5])}\n"
                        f"{r['summary']}\n"
                        f"Snippet: {r['snippet'][:400]}"
                    )
                    sources_used.append(r)

            context = "\n---\n".join(context_parts) if context_parts else "No directly relevant entries found."

            if HAS_OPENAI and os.getenv("OPENAI_API_KEY"):
                response = self._llm_chat(user_input, context, history)
            else:
                response = self._local_chat(user_input, results)

            console.print("\n[bold green]Cortex[/]")
            console.print(RichMarkdown(response))

            if sources_used:
                console.print(f"\n[dim]📚 Sources ({len(sources_used)}):[/]")
                for s in sources_used[:3]:
                    console.print(f"  [dim]• {s['title'][:60]} ({s['relevance']:.0%})[/]")

            history.append({"role": "user", "content": user_input})
            history.append({"role": "assistant", "content": response})
            now = datetime.now(timezone.utc).isoformat()
            self.db.execute("INSERT INTO chat_history (role,content,timestamp) VALUES (?,?,?)", ("user", user_input, now))
            self.db.execute(
                "INSERT INTO chat_history (role,content,timestamp) VALUES (?,?,?)",
                ("assistant", response, now),
            )
            self.db.commit()

    def _llm_chat(self, query: str, context: str, history: list) -> str:
        client = openai.OpenAI()
        system_prompt = (
            "You are The Cortex, a personal knowledge assistant. "
            "You have access to the user's curated knowledge base. "
            "Answer questions by synthesizing information from the provided context. "
            "Be precise, cite which sources you're drawing from, note connections between ideas, "
            "and flag when information might be incomplete or when the user should dig deeper. "
            "Be conversational but smart — the user appreciates depth and nuance."
        )
        messages = [{"role": "system", "content": system_prompt}]
        for h in history[-6:]:
            messages.append(h)
        messages.append(
            {
                "role": "user",
                "content": f"Knowledge Base Context:\n{context}\n\n---\nQuestion: {query}",
            }
        )
        try:
            resp = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=messages,
                max_tokens=800,
                temperature=0.7,
            )
            return resp.choices[0].message.content.strip()
        except Exception as e:
            return f"LLM error: {e}\n\nFalling back to local search results."

    def _local_chat(self, query: str, results: list) -> str:
        if not results:
            return (
                "I don't have anything on that yet. You could:\n"
                "- Ingest some relevant files or URLs\n"
                "- Run `/expand` to auto-discover sources\n"
                "- Ask me something else!"
            )
        response = "Here's what I found across your knowledge base:\n\n"
        for i, r in enumerate(results[:5], 1):
            response += f"**{i}. {r['title']}** ({r['quality']}, {r['relevance']:.0%} match)\n"
            response += f"   {r['summary'][:200]}\n"
            response += f"   Tags: {', '.join(r['tags'][:5])}\n\n"
        if len(results) > 1:
            all_tags = set()
            for r in results:
                all_tags.update(r["tags"])
            response += f"**Themes across results:** {', '.join(sorted(all_tags)[:8])}\n"
        return response

    def _handle_slash_command(self, cmd: str):
        parts = cmd.split(maxsplit=1)
        command = parts[0].lower()
        arg = parts[1] if len(parts) > 1 else ""

        if command == "/search":
            results = self.search(arg or "recent")
            for r in results[:5]:
                console.print(f"  • [cyan]{r['title']}[/] ({r['relevance']:.0%}) — {r['summary'][:80]}")
        elif command == "/expand":
            self.expand(arg or None)
        elif command == "/tags":
            tags = self.get_tags()
            for t, c in list(tags.items())[:15]:
                console.print(f"  🏷️  [cyan]{t}[/] ({c})")
        elif command == "/stats":
            s = self.get_stats()
            console.print(f"  📚 {s['total_entries']} entries · 🏷️ {s['unique_tags']} tags · 🔗 {s['relationships']} links")
        elif command == "/ingest":
            if arg:
                self.ingest(arg)
            else:
                console.print("[yellow]Usage: /ingest <file_or_url_or_text>[/]")
        elif command in ("/help", "/?"):
            console.print(
                "[dim]/search <query> — search knowledge base\n"
                "/expand [topic] — auto-discover new sources\n"
                "/ingest <source> — add knowledge inline\n"
                "/tags — view tag map\n"
                "/stats — overview\n"
                "/quit — exit chat[/]"
            )
        else:
            console.print(f"[yellow]Unknown command: {command}. Try /help[/]")

    def get_tags(self) -> dict[str, int]:
        rows = self.db.execute("SELECT tag, COUNT(*) as cnt FROM tags GROUP BY tag ORDER BY cnt DESC").fetchall()
        return {r["tag"]: r["cnt"] for r in rows}

    def get_stats(self) -> dict:
        total = self.db.execute("SELECT COUNT(*) as c FROM entries").fetchone()["c"]
        by_type = self.db.execute("SELECT source_type, COUNT(*) as c FROM entries GROUP BY source_type").fetchall()
        by_quality = self.db.execute("SELECT quality, COUNT(*) as c FROM entries GROUP BY quality").fetchall()
        tags = self.db.execute("SELECT COUNT(DISTINCT tag) as c FROM tags").fetchone()["c"]
        relations = self.db.execute("SELECT COUNT(*) as c FROM relationships").fetchone()["c"]
        return {
            "total_entries": total,
            "by_type": {r["source_type"]: r["c"] for r in by_type},
            "by_quality": {r["quality"]: r["c"] for r in by_quality},
            "unique_tags": tags,
            "relationships": relations,
        }

    def get_related(self, entry_id: str) -> list[dict]:
        rows = self.db.execute(
            """
            SELECT e.id, e.title, e.summary, r.relation, r.score
            FROM relationships r JOIN entries e ON e.id = r.target_id
            WHERE r.source_id = ? ORDER BY r.score DESC
            """,
            (entry_id,),
        ).fetchall()
        return [dict(r) for r in rows]

    def repair_malformed_auto_entries(self, reingest: bool = True, quiet: bool = False) -> dict:
        rows = self.db.execute(
            "SELECT id, source, title, summary, content FROM entries WHERE source_type = ?",
            (SourceType.AUTO.value,),
        ).fetchall()

        malformed = []
        for row in rows:
            source = (row["source"] or "").strip()
            title = (row["title"] or "").strip()
            summary = (row["summary"] or "").strip()
            content = (row["content"] or "").strip()
            url_like_content = content.startswith(("http://", "https://")) and len(content.split()) < 12
            is_malformed = source.startswith(("http://", "https://")) and (
                title == source or summary == source or url_like_content
            )
            if is_malformed:
                malformed.append((row["id"], source))

        removed = 0
        for entry_id, _ in malformed:
            self.db.execute("DELETE FROM tags WHERE entry_id = ?", (entry_id,))
            self.db.execute("DELETE FROM relationships WHERE source_id = ? OR target_id = ?", (entry_id, entry_id))
            self.db.execute("DELETE FROM entries WHERE id = ?", (entry_id,))
            if self.collection:
                try:
                    self.collection.delete(where={"entry_id": entry_id})
                except Exception:
                    pass
            removed += 1

        self.db.commit()

        reingested = 0
        if reingest:
            for _, url in malformed:
                entry = self.ingest(url, SourceType.AUTO, quality=QualityTier.SILVER, quiet=True)
                if entry:
                    reingested += 1

        if not quiet:
            console.print(
                Panel(
                    f"[bold green]✓ Repair complete[/]\n"
                    f"[dim]Malformed found:[/] {len(malformed)}\n"
                    f"[dim]Removed:[/] {removed}\n"
                    f"[dim]Reingested:[/] {reingested if reingest else 0}",
                    title="🛠️ Knowledge Repair",
                    border_style="green",
                )
            )

        return {
            "malformed_found": len(malformed),
            "removed": removed,
            "reingested": reingested if reingest else 0,
        }

    def run_daemon(self):
        if not HAS_SCHEDULE:
            console.print("[red]Install schedule: pip install schedule[/]")
            return

        console.print(
            Panel(
                f"[bold magenta]🤖 Cortex Daemon Running[/]\n\n"
                f"[dim]Feed pull every {FEED_PULL_INTERVAL}h\n"
                f"Knowledge expansion every {EXPAND_INTERVAL}h\n\n"
                f"Press Ctrl+C to stop.[/]",
                border_style="magenta",
            )
        )

        def pull_job():
            console.print(f"\n[dim][{datetime.now().strftime('%H:%M')}] Running feed pull...[/]")
            try:
                self.feed_pull(quiet=True)
            except Exception as e:
                console.print(f"[red]Feed pull error: {e}[/]")

        def expand_job():
            console.print(f"\n[dim][{datetime.now().strftime('%H:%M')}] Running auto-expansion...[/]")
            try:
                self.expand(quiet=True)
            except Exception as e:
                console.print(f"[red]Expansion error: {e}[/]")

        schedule.every(FEED_PULL_INTERVAL).hours.do(pull_job)
        schedule.every(EXPAND_INTERVAL).hours.do(expand_job)

        pull_job()

        try:
            while True:
                schedule.run_pending()
                time.sleep(60)
        except KeyboardInterrupt:
            console.print("\n[dim]Daemon stopped. Knowledge persists. 🧠[/]")


@app.command()
def ingest(source: str = typer.Argument(..., help="File path, URL, or text note")):
    """Ingest a file, URL, or note into the knowledge base."""
    cortex = Cortex()
    cortex.ingest(source)


@app.command()
def search(query: str = typer.Argument(..., help="Natural language search query")):
    """Search your knowledge base."""
    cortex = Cortex()
    results = cortex.search(query)
    if not results:
        console.print("[yellow]No results found.[/]")
        return
    for i, r in enumerate(results, 1):
        quality_color = {"gold": "yellow", "silver": "white", "bronze": "red"}.get(r["quality"], "dim")
        table = Table(box=box.ROUNDED, show_header=False, border_style="cyan", width=90)
        table.add_column(width=15)
        table.add_column(width=72)
        table.add_row("[bold]Title[/]", r["title"][:70])
        table.add_row("[bold]Relevance[/]", f"{'█' * int(r['relevance'] * 20)} {r['relevance']:.0%}")
        table.add_row("[bold]Quality[/]", f"[{quality_color}]{r['quality']}[/]")
        table.add_row("[bold]Tags[/]", ", ".join(r["tags"][:6]))
        table.add_row("[bold]Summary[/]", textwrap.shorten(r["summary"], 200))
        table.add_row("[bold]Source[/]", textwrap.shorten(r["source"], 70))
        console.print(f"\n[bold cyan]#{i}[/]")
        console.print(table)


@app.command()
def expand(topic: str = typer.Argument(None, help="Optional topic to expand on")):
    """Auto-discover and ingest new relevant sources."""
    cortex = Cortex()
    cortex.expand(topic)


@app.command()
def explore(topic: str = typer.Argument(..., help="Topic to deep-dive")):
    """Deep exploration: search + expand combined."""
    cortex = Cortex()
    console.print(Panel(f"🧠 [bold]Deep-diving: {topic}[/]", border_style="magenta"))
    results = cortex.search(topic)
    if results:
        console.print(f"\n[bold]Found {len(results)} existing entries:[/]")
        for r in results[:3]:
            console.print(f"  • [cyan]{r['title']}[/] ({r['relevance']:.0%})")
    cortex.expand(topic)
    console.print(f"\n[bold]Updated knowledge on '{topic}':[/]")
    results = cortex.search(topic)
    for r in results[:5]:
        console.print(f"  • [{r['quality']}] [cyan]{r['title']}[/] — {r['summary'][:100]}")


@app.command()
def tags():
    """View the knowledge map by tags."""
    cortex = Cortex()
    tag_counts = cortex.get_tags()
    if not tag_counts:
        console.print("[yellow]No tags yet.[/]")
        return
    table = Table(title="🏷️  Knowledge Map", box=box.DOUBLE_EDGE, border_style="cyan")
    table.add_column("Tag", style="bold cyan")
    table.add_column("Entries", justify="right")
    table.add_column("Bar", width=30)
    max_count = max(tag_counts.values())
    for tag, count in tag_counts.items():
        bar = "█" * int(count / max_count * 28)
        table.add_row(tag, str(count), f"[green]{bar}[/]")
    console.print(table)


@app.command()
def stats():
    """Show knowledge base overview."""
    cortex = Cortex()
    s = cortex.get_stats()
    panel_text = (
        f"[bold]📚 Total Entries:[/]    {s['total_entries']}\n"
        f"[bold]🏷️  Unique Tags:[/]     {s['unique_tags']}\n"
        f"[bold]🔗 Relationships:[/]   {s['relationships']}\n\n"
        f"[bold]By Type:[/]\n"
    )
    for t, c in s.get("by_type", {}).items():
        panel_text += f"  {t}: {c}\n"
    panel_text += "\n[bold]By Quality:[/]\n"
    for q, c in s.get("by_quality", {}).items():
        emoji = {"gold": "🥇", "silver": "🥈", "bronze": "🥉"}.get(q, "❓")
        panel_text += f"  {emoji} {q}: {c}\n"
    console.print(Panel(panel_text, title="📊 Cortex Stats", border_style="cyan"))


@app.command()
def related(entry_id: str = typer.Argument(..., help="Entry ID")):
    """Show entries related to a specific entry."""
    cortex = Cortex()
    rels = cortex.get_related(entry_id)
    if not rels:
        console.print("[yellow]No relationships found.[/]")
        return
    for r in rels:
        console.print(f"  [{r['relation']}] ({r['score']:.2f}) → [cyan]{r['title']}[/]")
        console.print(f"    {r['summary'][:120]}\n")


@app.command(name="chat")
def chat_cmd():
    """Interactive conversational interface to your knowledge base."""
    cortex = Cortex()
    cortex.chat()


@app.command()
def daemon():
    """Run scheduled feed pulls and auto-expansion in the background."""
    cortex = Cortex()
    cortex.run_daemon()


@app.command()
def repair(
    reingest: bool = typer.Option(
        True,
        "--reingest/--no-reingest",
        help="Re-ingest removed malformed URLs using the fixed pipeline.",
    ),
    quiet: bool = typer.Option(
        False,
        "--quiet",
        help="Suppress rich panel output.",
    ),
    json_output: bool = typer.Option(
        False,
        "--json",
        help="Emit machine-readable JSON output.",
    ),
):
    """Repair malformed auto-discovered URL-only entries."""
    cortex = Cortex()
    result = cortex.repair_malformed_auto_entries(reingest=reingest, quiet=quiet or json_output)
    if json_output:
        console.print(json.dumps(result, indent=2))


@feeds_app.command("add")
def feed_add(
    url: str = typer.Argument(..., help="RSS/Atom feed URL"),
    tag: str = typer.Option("", help="Tag to apply to ingested articles"),
    name: str = typer.Option(None, help="Custom name for this feed"),
):
    """Subscribe to an RSS/Atom feed."""
    cortex = Cortex()
    cortex.feed_add(url, name=name, tag=tag)


@feeds_app.command("list")
def feed_list():
    """List all subscribed feeds."""
    cortex = Cortex()
    cortex.feed_list()


@feeds_app.command("pull")
def feed_pull():
    """Pull new articles from all active feeds now."""
    cortex = Cortex()
    cortex.feed_pull()


@feeds_app.command("remove")
def feed_remove(feed_id: str = typer.Argument(..., help="Feed ID to deactivate")):
    """Remove (deactivate) a feed subscription."""
    cortex = Cortex()
    cortex.feed_remove(feed_id)


@export_app.command("obsidian")
def export_obsidian(dir: Path = typer.Option(None, "--dir", help="Target vault directory")):
    """Export knowledge base as an Obsidian vault."""
    cortex = Cortex()
    cortex.export_obsidian(dir)


@export_app.command("json")
def export_json(file: Path = typer.Option(None, "--file", help="Output JSON file")):
    """Export knowledge base as structured JSON."""
    cortex = Cortex()
    cortex.export_json(file)


if __name__ == "__main__":
    app()